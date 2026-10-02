from __future__ import annotations

import io

import openpyxl

from .errors import BackendError

SHEET_NAME = "Лист1"
# Колонки A, B, C — место, команда, итого; раунды идут дальше, начиная с D.
FIXED_COLUMNS = 3


def read_teams(data: bytes, rounds: int = 0) -> list[dict]:
    """Читает таблицу рейтинга из байтов .xlsx и возвращает список команд.

    Каждая запись: {"place": int|str, "name": str, "total": Any, "rounds": [rounds значений]}.
    Данные читаются со 2-й строки листа SHEET_NAME до первой строки без места или названия.
    Из колонок раундов (с D) берётся ровно `rounds` штук: сколько их — решает таблица в Keynote,
    а не Excel. Пустая или отсутствующая ячейка даёт None.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise BackendError("EXCEL_BAD_FORMAT", f"Не удалось открыть таблицу: {exc}") from exc

    try:
        if SHEET_NAME not in wb.sheetnames:
            raise BackendError(
                "EXCEL_BAD_FORMAT",
                f"Лист '{SHEET_NAME}' не найден. Доступные листы: {', '.join(wb.sheetnames)}",
            )

        width = FIXED_COLUMNS + rounds
        teams: list[dict] = []
        for row in wb[SHEET_NAME].iter_rows(min_row=2, values_only=True):
            row = tuple(row) + (None,) * max(0, width - len(row))
            place, name, total = row[0], row[1], row[2]  # A: место, B: команда, C: итого

            if place is None or name is None:
                break

            teams.append({
                "place": int(place) if isinstance(place, float) and place == int(place) else place,
                "name": str(name),
                "total": total,
                "rounds": list(row[FIXED_COLUMNS:width]),
            })
        return teams
    finally:
        wb.close()


def teams_by_place(teams: list[dict]) -> dict[int, str]:
    """Строит словарь {место: имя команды} из результата read_teams()."""
    result: dict[int, str] = {}
    for team in teams:
        place = team["place"]
        place_int = int(place)
        result[place_int] = team["name"]
    return result
